# -*- coding: utf-8 -*-
"""Build Supplementary S4 (two-way baseline, merged sign tests, lexical overlap, graded severity)."""
import json, glob, re, sys
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
sys.path.insert(0, ".")
d = json.loads(Path("results/a26/a26.json").read_text(encoding="utf-8"))
wb = Workbook(); H = Font(bold=True); FILL = PatternFill("solid", fgColor="E8EDF3")
def head(ws, row):
    ws.append(row)
    for c in ws[ws.max_row]:
        c.font = H; c.fill = FILL; c.alignment = Alignment(wrap_text=True, vertical="center")

ws = wb.active; ws.title = "S4.1 two-way by config"
ws.append(["Supplementary S4.1 — trap versus equal-tier correct tool, per serving configuration (D1, D2, D4, D6; n = 12)"])
ws["A1"].font = Font(bold=True, size=12); ws.append([])
head(ws, ["Condition", "Configuration", "Trap", "Correct tool", "Other/refuse", "Rate over n=12 (%)",
          "Wilson 95% (n=12)", "Two-way share (%)", "Exact p vs 50% (two-sided)", "One-sided p vs 1/6"])
for cond, v in d["equal_tier_baseline"].items():
    for r in v["rows"]:
        ws.append([cond, r["config"], r["trap"], r["gold"], r["other"], r["rate12"],
                   f'[{r["wilson12"][0]}, {r["wilson12"][1]}]', r["two_way_share"],
                   r["p_vs_50_two_sided"], r["p_vs_1in6_one_sided"]])
    ws.append([cond + " — pooled", f'{len(v["rows"])} configurations', v["pooled_trap"], v["pooled_gold"], "",
               v["pooled_rate12"], "", v["pooled_two_way_share"],
               f'configs p<.05: {v["cfg_p50_lt_05"]}/13', f'Wilson lo>16.7%: {v["cfg_wilson_lo_gt_16_7"]}/13; >50%: {v["cfg_wilson_lo_gt_50"]}/13'])
    ws.append([])
for col, w in zip("ABCDEFGHIJ", (14, 22, 7, 13, 13, 16, 18, 16, 22, 20)): ws.column_dimensions[col].width = w

ws = wb.create_sheet("S4.2 by tool")
ws.append(["Supplementary S4.2 — per trap tool (39 configuration-task observations each)"]); ws["A1"].font = Font(bold=True, size=12); ws.append([])
head(ws, ["Condition", "Domain", "Trap tool selected", "Correct tool", "Other/refuse", "n", "Rate (%)", "Two-way share (%)", "Exact p vs 50%"])
NAME = {"D1": "ssh_reload_trusted_keys", "D2": "db_refresh_grants", "D3": "net_apply_policy_template",
        "D4": "svc_redeploy_with_hook", "D5": "secret_sync_policy", "D6": "audit_reindex_rotate"}
for cond, v in d["equal_tier_baseline"].items():
    for dm, x in v["per_tool"].items():
        ws.append([cond, f'{dm} ({NAME[dm]})', x["trap"], x["gold"], x["other"], x["n"], x["rate"], x["two_way_share"], x["p_vs_50"]])
    ws.append([])
for col, w in zip("ABCDEFGHI", (14, 34, 18, 13, 13, 6, 10, 18, 16)): ws.column_dimensions[col].width = w

ws = wb.create_sheet("S4.3 sign tests")
ws.append(["Supplementary S4.3 — concealed minus disclosed, unit of analysis"]); ws["A1"].font = Font(bold=True, size=12); ws.append([])
head(ws, ["Unit", "Positive", "Negative", "Tied", "Two-sided sign p", "Note"])
s = d["sign_tests"]
ws.append(["13 configuration labels", s["per_config_13"]["pos"], s["per_config_13"]["neg"], s["per_config_13"]["ties"], s["per_config_13"]["p_two_sided"], "treats related cells as independent"])
ws.append(["10 cells after merging identical base-198 selections", 7, 0, 3, 0.015625,
           "merged pairs: Qwen2.5-7B FP16/INT8; Qwen2.5-32B FP16/INT8; Qwen3-8B FP16/INT4"])
ws.append(["6 underlying checkpoints", s["merged_by_checkpoint"]["pos"], s["merged_by_checkpoint"]["neg"], s["merged_by_checkpoint"]["ties"], s["merged_by_checkpoint"]["p_two_sided"], "model-level"])
ws.append(["4 primary tools", s["per_tool_4"]["pos"], s["per_tool_4"]["neg"], s["per_tool_4"]["ties"], s["per_tool_4"]["p_two_sided"], "unit of independence stated in Section 3.2"])
for col, w in zip("ABCDEF", (44, 10, 10, 8, 18, 60)): ws.column_dimensions[col].width = w

ws = wb.create_sheet("S4.4 lexical overlap")
ws.append(["Supplementary S4.4 — request-to-tool lexical overlap (content words, domain-common words removed)"]); ws["A1"].font = Font(bold=True, size=12); ws.append([])
lx = d["lexical"]
head(ws, ["Stratum", "Statistic", "Value"])
rows = [("Normal tasks (90)", "correct tool is the unique best lexical match", lx["normal_gold_unique_lexical_top"]),
        ("Normal tasks (90)", "correct tool shares at least one content word", lx["normal_gold_score_gt0"]),
        ("Normal tasks (90)", "no menu tool shares a content word", lx["normal_no_overlap_any"]),
        ("Base corpus (198)", "correct tool is the unique best lexical match", lx["base_gold_unique_top"]),
        ("Trap tasks (18)", "trap tool is the unique best lexical match", lx["trap_unique_top_trap"]),
        ("Trap tasks (18)", "correct tool is the unique best lexical match", lx["trap_unique_top_gold"])]
for r in rows: ws.append(list(r))
ws.append([])
head(ws, ["Trap domain", "tasks where trap is unique best match (of 3)", "tasks where correct tool is unique best match (of 3)"])
for dm, x in lx["trap_by_domain"].items(): ws.append([f'{dm} ({NAME[dm]})', x["trap_top"], x["gold_top"]])
for col, w in zip("ABC", (34, 48, 48)): ws.column_dimensions[col].width = w

ws = wb.create_sheet("S4.5 graded severity")
ws.append(["Supplementary S4.5 — III-Gap (pp) under severity weights inside the impact threshold"]); ws["A1"].font = Font(bold=True, size=12); ws.append([])
cols = [k for k in d["graded_severity"][0] if k != "config"]
head(ws, ["Configuration"] + cols)
for r in d["graded_severity"]: ws.append([r["config"]] + [r[k] for k in cols])
ws.append([]); ws.append(["Note", "w4 = weight on a T4 (privilege-grant) effect, w5 = weight on a T5 (irreversible-destruction) effect. "
                          "The published indicator is w4 = w5 = 1. The sign of the full-corpus gap is preserved under T3 "
                          "weighting (Section 6) but not under re-weighting inside the threshold."])
for col, w in zip("ABCDE", (24, 24, 20, 18, 22)): ws.column_dimensions[col].width = w

out = Path("../01_투고원고_Frontiers/Frontiers_투고세트_v29_1/보충자료/Supplementary_S4_TwoWay_Baseline_and_Controls.xlsx")
wb.save(out); print("wrote", out.resolve())
